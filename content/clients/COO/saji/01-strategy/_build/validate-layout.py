"""Geometry check: no text box may overlap another, overflow its own shape,
or cross the footer bar. Run after every build."""
import sys, math, re
from pptx import Presentation
from pptx.util import Emu

E = 914400.0
def est_h(txt, w_in, pt, ls=1.10):
    cpl = max(6.0, w_in * (17.6*8.5/pt) * 0.84)
    lines = sum(max(1, math.ceil(len(p)/cpl)) for p in str(txt).split("\n"))
    return lines * pt * ls / 72.0

prs = Presentation(sys.argv[1] if len(sys.argv)>1 else 'rev3-cleaned.pptx')
problems = []

for n, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text.strip()
        if not t: continue
        L,T = sh.left/E, sh.top/E
        W,H = sh.width/E, sh.height/E
        # biggest font in the shape drives the estimate
        pt = max([r.font.size.pt for p in sh.text_frame.paragraphs
                  for r in p.runs if r.font.size] or [10])
        ml = (sh.text_frame.margin_left or 0)/E
        mr = (sh.text_frame.margin_right or 0)/E
        mt = (sh.text_frame.margin_top or 0)/E
        need = sum(est_h(p.text, W-ml-mr, max([r.font.size.pt for r in p.runs if r.font.size] or [pt]),
                         p.line_spacing if isinstance(p.line_spacing,float) else 1.10)
                   + (p.space_after.pt/72.0 if p.space_after else 0)
                   for p in sh.text_frame.paragraphs if p.text.strip())
        boxes.append((L,T,W,H,t[:44],need+mt))

    # 1. text overflowing its own shape
    for L,T,W,H,t,need in boxes:
        if need > H + 0.02:
            problems.append(f"slide {n}: text needs {need:.2f}in in a {H:.2f}in box | {t}")
    # 2. text boxes overlapping each other
    for i in range(len(boxes)):
        for j in range(i+1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ox = min(a[0]+a[2], b[0]+b[2]) - max(a[0], b[0])
            oy = min(a[1]+a[3], b[1]+b[3]) - max(a[1], b[1])
            if ox > 0.02 and oy > 0.02:
                problems.append(f"slide {n}: OVERLAP {ox:.2f}x{oy:.2f}in | '{a[4]}' vs '{b[4]}'")
    # 3. content crossing into the footer bar, on slides that have one
    has_bar = any(abs(sh.top/E - 4.86) < 0.02 and abs(sh.height/E - 0.44) < 0.02
                  for sh in slide.shapes)
    if has_bar:
        for L,T,W,H,t,need in boxes:
            if T < 4.84 and T + max(H, need) > 4.86:
                problems.append(f"slide {n}: crosses the footer bar | {t}")

print("PROBLEMS:", len(problems))
for p in problems: print("  " + p)
if not problems: print("  clean: no overflow, no overlap, no footer collision")
